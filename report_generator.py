
from docx import Document
from docx.shared import Inches, Pt
from io import BytesIO
from datetime import datetime


def generate_report(result, file_a, file_b):
    doc = Document()

    section = doc.sections[0]
    section.header.paragraphs[0].text = (
        "AI DOCUMENT COMPARISON ASSISTANT"
    )

    doc.add_heading("AI Document Comparison Report", 0)

    doc.add_paragraph(
        "Generated on: " +
        datetime.now().strftime("%d %B %Y, %I:%M %p")
    )

    doc.add_heading("Documents Compared", level=1)
    doc.add_paragraph("Document A: " + file_a)
    doc.add_paragraph("Document B: " + file_b)

    doc.add_heading("AI Comparison Results", level=1)

    for line in result.split("\n"):
        line = line.strip()

        if not line:
            continue

        if line.startswith("## "):
            doc.add_heading(line.replace("## ", ""), level=2)

        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")

        else:
            doc.add_paragraph(line)

    doc.add_heading("Disclaimer", level=1)
    doc.add_paragraph(
        "This report is generated using AI and is intended "
        "for informational purposes. Verify important "
        "details against the original documents."
    )

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output.getvalue()
